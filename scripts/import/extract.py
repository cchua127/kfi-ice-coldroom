#!/usr/bin/env python3
"""
Stage 1 of the migration: dump every cell of every source workbook to JSON.

This tool does NO interpretation. It records what each cell holds — cached
value, formula text, and whether Excel stored it as a date — and nothing else.
All business logic lives in TypeScript, under test, in src/lib/import.

The split exists for two reasons. Two of the six workbooks are legacy BIFF
(.xls), which no maintained JavaScript library reads, and the one that does
carries unpatched advisories. And a dumb, reviewable intermediate makes the
migration auditable: you can diff the extract itself before anything is loaded.

Usage:
    python3 scripts/import/extract.py <source-dir> <out-dir>

Requires: openpyxl, xlrd  (pip install openpyxl xlrd)
"""
from __future__ import annotations
import json, sys, os, datetime, hashlib

WORKBOOKS = {
    "daily-rekod-ais":  ["Daily rekod Ais", "Daily_rekod_Ais"],
    "meter-tube":       ["Meter - tube ice - elec", "Meter_-_tube_ice_-_elec"],
    "meter-big-pool":   ["meter BIG POOL vs elec", "meter_BIG_POOL_vs_elec"],
    "meter-small-pool": ["meter SMALL POOL vs elec", "meter_SMALL_POOL_vs_elec"],
    "daily-cash-ice":   ["daily cash ice", "daily_cash_ice"],
    "ice-purchase":     ["Ice Purchase  sales outside", "Ice_Purchase__sales_outside"],
}


def cell_record(value, formula=None, is_date=False):
    """One cell, as it actually sits in the file."""
    if isinstance(value, (datetime.datetime, datetime.date)):
        return {"t": "date", "v": value.isoformat()[:10], "f": formula}
    if isinstance(value, bool):
        return {"t": "bool", "v": value, "f": formula}
    if isinstance(value, (int, float)):
        return {"t": "date" if is_date else "num", "v": value, "f": formula}
    if value is None:
        return None
    return {"t": "str", "v": str(value), "f": formula}


def read_xlsx(path):
    import openpyxl
    wf = openpyxl.load_workbook(path, data_only=False)
    wv = openpyxl.load_workbook(path, data_only=True)
    out = {}
    for name in wf.sheetnames:
        sf, sv = wf[name], wv[name]
        rows = {}
        for r in range(1, sf.max_row + 1):
            for c in range(1, sf.max_column + 1):
                raw, val = sf.cell(r, c).value, sv.cell(r, c).value
                formula = raw if isinstance(raw, str) and raw.startswith("=") else None
                rec = cell_record(val, formula)
                if rec is None and formula is None:
                    continue
                rows.setdefault(str(r), {})[
                    openpyxl.utils.get_column_letter(c)
                ] = rec if rec else {"t": "empty", "v": None, "f": formula}
        out[name] = {"maxRow": sf.max_row, "maxCol": sf.max_column, "cells": rows}
    return out


def read_xls(path):
    import xlrd
    wb = xlrd.open_workbook(path)
    out = {}
    for sh in wb.sheets():
        rows = {}
        for r in range(sh.nrows):
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype in (0, 6):  # empty / blank
                    continue
                if cell.ctype == 3:  # Excel serial that Excel typed as a date
                    try:
                        tup = xlrd.xldate_as_tuple(cell.value, wb.datemode)
                        v = datetime.datetime(*tup)
                        rec = {"t": "date", "v": v.isoformat()[:10], "f": None,
                               "serial": cell.value}
                    except Exception:
                        rec = {"t": "num", "v": cell.value, "f": None}
                else:
                    rec = cell_record(cell.value)
                if rec is None:
                    continue
                rows.setdefault(str(r + 1), {})[xlrd.colname(c)] = rec
        out[sh.name] = {"maxRow": sh.nrows, "maxCol": sh.ncols, "cells": rows}
    return out


def find(src_dir, patterns):
    for fn in sorted(os.listdir(src_dir)):
        low = fn.lower()
        if not (low.endswith(".xlsx") or low.endswith(".xls")):
            continue
        for p in patterns:
            if p.lower().replace(" ", "_") in low.replace(" ", "_"):
                return os.path.join(src_dir, fn)
    return None


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    src_dir, out_dir = sys.argv[1], sys.argv[2]
    os.makedirs(out_dir, exist_ok=True)

    manifest = []
    for key, patterns in WORKBOOKS.items():
        path = find(src_dir, patterns)
        if not path:
            print(f"  MISSING  {key}")
            manifest.append({"key": key, "found": False})
            continue
        sha = hashlib.sha256(open(path, "rb").read()).hexdigest()
        sheets = read_xls(path) if path.lower().endswith(".xls") else read_xlsx(path)
        payload = {
            "workbook": key,
            "sourceFile": os.path.basename(path),
            "sha256": sha,
            "extractedAt": datetime.datetime.now(datetime.timezone.utc).isoformat(),
            "sheets": sheets,
        }
        dest = os.path.join(out_dir, f"{key}.json")
        with open(dest, "w") as fh:
            json.dump(payload, fh, indent=1, sort_keys=True)
        print(f"  ok  {key:18s} {len(sheets):2d} sheets  <- {os.path.basename(path)}")
        manifest.append({
            "key": key, "found": True, "sourceFile": os.path.basename(path),
            "sha256": sha, "sheets": sorted(sheets.keys()),
        })

    with open(os.path.join(out_dir, "manifest.json"), "w") as fh:
        json.dump(manifest, fh, indent=1)
    missing = [m["key"] for m in manifest if not m["found"]]
    print(f"\n{len(manifest) - len(missing)}/{len(manifest)} workbooks extracted"
          + (f"; missing: {', '.join(missing)}" if missing else ""))


if __name__ == "__main__":
    main()
